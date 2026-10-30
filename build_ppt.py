"""
Seva Cyber Sandbox — Automated PPT Generator
Uses python-pptx to inject content and screenshots into the official template.
"""

import os
from pptx import Presentation
from pptx.util import Inches

TEMPLATE_PATH = "Seva-First-Innovation-Challenge-PPT-Template.pptx"
OUTPUT_PATH = "Seva_Cyber_Sandbox_Submission.pptx"

def main():
    if not os.path.exists(TEMPLATE_PATH):
        print(f"Error: {TEMPLATE_PATH} not found in this folder.")
        return

    # Load the official template
    prs = Presentation(TEMPLATE_PATH)

    # ---------------------------------------------------------
    # SLIDE 1: Title & Applicant Snapshot
    # ---------------------------------------------------------
    slide1 = prs.slides[0]
    
    # We will just write to the first two shapes assuming they are title/subtitle placeholders
    if len(slide1.shapes) >= 2:
        slide1.shapes[0].text = "Seva Cyber Sandbox (Maha-Suraksha)"
        slide1.shapes[1].text = (
            "A privacy-preserving static Android .apk and Marathi/mixed-language "
            "WhatsApp phishing sandbox that warns citizens of scam indicators in "
            "under 2 seconds before they install untrusted apps or share OTPs."
        )

    # Add Speaker Notes
    slide1.notes_slide.notes_text_frame.text = (
        "Namaskar to the Grand Jury. I am Pratik Chendake from AIT Pune. Every day in Maharashtra, "
        "smartphone users receive urgent WhatsApp messages in Marathi threatening electricity "
        "disconnection at 9:30 PM or promising welfare payouts, attached with a raw Android .apk file. "
        "Installing that file lets scammers silently intercept bank SMS OTPs. Seva Cyber Sandbox is a "
        "live, working defensive prototype that inspects suspicious .apk files and Marathi messages "
        "statically in under 2 seconds—explaining the exact risk in plain Marathi and English before "
        "financial loss occurs."
    )

    # ---------------------------------------------------------
    # SLIDE 2: The Problem
    # ---------------------------------------------------------
    slide2 = prs.slides[1]
    slide2.notes_slide.notes_text_frame.text = (
        "Following Step 1 of the Seva Problem Solver's Guide, we started from a real problem observed "
        "around us in Pune. Scammers no longer rely on English emails; they send urgent Marathi WhatsApp "
        "messages threatening tonight's electricity cut alongside an .apk file. In our local user "
        "consultation, 5 out of 5 smartphone users had zero way to inspect what permissions an .apk "
        "file secretly asks for before tapping install."
    )

    # ---------------------------------------------------------
    # SLIDE 3: Our Solution & Innovation
    # ---------------------------------------------------------
    slide3 = prs.slides[2]
    
    # Insert Screenshot 1 & 2 on the right side
    img1_path = "screenshot_1_phishing.png"
    img2_path = "screenshot_2_safeguard.png"
    
    if os.path.exists(img1_path):
        # Adjust dimensions as needed for your specific template layout
        slide3.shapes.add_picture(img1_path, Inches(6.5), Inches(1.5), width=Inches(3.0))
    if os.path.exists(img2_path):
        slide3.shapes.add_picture(img2_path, Inches(6.5), Inches(4.0), width=Inches(3.0))

    slide3.notes_slide.notes_text_frame.text = (
        "Slide 3 shows our live prototype. On the right, you see two real screenshots from our deployed "
        "app: first, a Marathi MSEDCL disconnection scam with an APK lure is flagged at 79 out of 100 "
        "with a full Marathi explanation. Second, right below it, a legitimate Marathi MSEDCL bill "
        "reminder scores 0 out of 100 with our green Context Safeguard active. Sharing that early "
        "failure—where keyword matching flagged normal reminders—and how we engineered context "
        "suppression is a core part of our methodology."
    )

    # ---------------------------------------------------------
    # SLIDE 4: Feasibility & Implementation
    # ---------------------------------------------------------
    slide4 = prs.slides[3]
    
    # Insert Screenshot 3
    img3_path = "screenshot_3_apk.png"
    if os.path.exists(img3_path):
        slide4.shapes.add_picture(img3_path, Inches(6.0), Inches(2.0), width=Inches(3.5))

    slide4.notes_slide.notes_text_frame.text = (
        "We are at the Working Prototype stage today, backed by 16 passing automated pytest tests and "
        "a live deployment. As shown in our third screenshot, our APK engine extracts critical "
        "permissions like BIND_ACCESSIBILITY_SERVICE and READ_SMS without ever running the file. Stage "
        "1 and 2 cost ₹0 in software licenses, and a 100-center Stage 3 pilot costs an estimated ₹2.4 "
        "Lakh—fitting comfortably inside existing cyber-awareness budgets and the ₹10 Lakh work order ceiling."
    )

    # ---------------------------------------------------------
    # SLIDE 5: Impact & Scalability
    # ---------------------------------------------------------
    slide5 = prs.slides[4]
    slide5.notes_slide.notes_text_frame.text = (
        "On Slide 5, we strictly separate measured baseline results from planned targets, just as the "
        "Seva Guide requires. In our 5-user prototype walkthrough, comprehension of hidden APK "
        "permissions went from 0 out of 5 before using the tool to 5 out of 5 after seeing the Marathi "
        "risk breakdown. To scale in circles without hiring new government staff, each site in Stages "
        "2 and 3 is run by people who already exist—NSS volunteers and Setu Kendra operators—using a "
        "one-page operating sheet."
    )

    # ---------------------------------------------------------
    # SLIDE 6: Team/Individual & Support Required
    # ---------------------------------------------------------
    slide6 = prs.slides[5]
    slide6.notes_slide.notes_text_frame.text = (
        "I am submitting as an individual student innovator from AIT Pune. Every line of code, the 16 "
        "automated tests, and the live Streamlit deployment are built and verifiable today. With "
        "mentorship, testing support, and a departmental pilot through the Seva First Innovation "
        "Challenge, I am ready to connect this verified sandbox to a public WhatsApp helpline and pilot "
        "it across citizen facilitation counters in Pune. Thank you!"
    )

    # Save the new presentation
    prs.save(OUTPUT_PATH)
    print(f"Successfully generated: {OUTPUT_PATH}")
    print("Open the file in PowerPoint to finalize formatting and text placement, then export to PDF.")

if __name__ == "__main__":
    main()